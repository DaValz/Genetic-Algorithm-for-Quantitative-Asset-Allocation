"""Test dell'esportazione Excel (nella lingua attiva) e JSON (chiavi in inglese, indipendente dalla lingua)."""
import io
import json

import numpy as np
import pytest
from openpyxl import load_workbook

from goalga import GAConfig, GoalProblem, i18n, run_ga
from goalga.analytics import build_report, run_robustness
from goalga.i18n import t
from goalga.plan import Plan
from ui import charts
from ui.export import build_excel, build_json

SHEETS = ["summary", "glide", "composition", "portfolios", "assumptions", "wealth", "goals", "comparison", "robustness"]


@pytest.fixture(scope="module")
def artefacts(market):
    plan = Plan(w0=60_000, contribution=4_000, horizon=15, goals=((8, 30_000.0), (15, 200_000.0)),
                pr_target=0.9, alpha=0.05)
    flows = plan.flows()
    problem = GoalProblem(flows, market.draw_scenarios(15, 800, np.random.default_rng(1)), 0.9, 0.05)
    result = run_ga(problem, GAConfig(pop_size=40, patience=20, seed=2))
    report = build_report(market, flows, problem, result, oos_scenarios=2000)
    rob = run_robustness(market, flows, pr_target=0.9, alpha=0.05, ga_config=GAConfig(pop_size=30, patience=10),
                         n_runs=3, n_scenarios=300)
    return plan, report, rob


def test_excel_workbook_structure(market, artefacts):
    plan, report, rob = artefacts
    data = build_excel(report, plan, market, {"seed": 2, "scenarios": 800}, rob)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Sintesi", "Glide path", "Composizione", "Portafogli", "Ipotesi", "Patrimonio",
                             "Obiettivi", "Confronto", "Robustezza"]
    assert wb["Glide path"].max_row == 1 + plan.horizon
    assert wb["Composizione"].max_column == 1 + market.n_assets
    assert wb["Patrimonio"].max_row == 2 + plan.horizon
    summary = {row[0].value: row[1].value for row in wb["Sintesi"].iter_rows(min_row=2)}
    assert summary[t("export.horizon")] == 15
    assert summary[t("export.pr_in")] == pytest.approx(report.pr)
    assert summary[t("export.setting.seed")] == 2 and summary[t("export.setting.scenarios")] == 800


def test_excel_without_robustness_has_no_extra_sheet(market, artefacts):
    plan, report, _ = artefacts
    wb = load_workbook(io.BytesIO(build_excel(report, plan, market, {}, None)))
    assert "Robustezza" not in wb.sheetnames


def test_percent_columns_are_formatted_as_percentages(market, artefacts):
    plan, report, _ = artefacts
    wb = load_workbook(io.BytesIO(build_excel(report, plan, market, {}, None)))
    glide = wb["Glide path"]
    headers = [c.value for c in glide[1]]
    col = headers.index(t("export.col_exp_ret")) + 1
    assert glide.cell(row=2, column=col).number_format == "0.00%"
    comparison = wb["Confronto"]
    prob = [c.value for c in comparison[1]].index(t("export.col_prob")) + 1
    assert comparison.cell(row=2, column=prob).number_format == "0.00%"


@pytest.mark.parametrize("lang", list(i18n.LANGS))
def test_excel_is_localized_in_every_language(market, artefacts, lang):
    """Fogli e intestazioni seguono la lingua attiva; il foglio dei confronti non duplica la colonna «Strategia»."""
    plan, report, rob = artefacts
    with i18n.use(lang):
        data = build_excel(report, plan, market, {"seed": 2}, rob)
        expected = [t(f"export.sheet.{key}") for key in SHEETS]
        strategy_header = t("export.col_strategy")
        labels = [charts.strategy_label(k) for k in report.benchmarks["Chiave"]]
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == expected
    assert all(len(name) <= 31 and not set("[]:*?/\\") & set(name) for name in wb.sheetnames)
    comparison = wb[expected[SHEETS.index("comparison")]]
    headers = [c.value for c in comparison[1]]
    assert headers[0] == strategy_header and len(set(headers)) == len(headers)
    assert [comparison.cell(row=r, column=1).value for r in range(2, 2 + len(labels))] == labels


def test_json_contains_everything_needed_to_reproduce(market, artefacts):
    plan, report, _ = artefacts
    payload = json.loads(build_json(report, plan, market, {"seed": 2}))
    assert payload["plan"]["horizon"] == 15 and payload["plan"]["goals"] == [[8, 30000.0], [15, 200000.0]]
    assert payload["settings"]["seed"] == 2
    assert payload["market"]["assets"] == market.names and len(payload["market"]["portfolio_weights"]) == 9
    assert payload["result"]["glide_path"] == report.path.tolist()
    assert payload["result"]["success_probability"] == pytest.approx(report.pr)


def test_json_says_which_probability_is_which_and_carries_the_validation_on_new_scenarios(market, artefacts):
    plan, report, _ = artefacts
    result = json.loads(build_json(report, plan, market, {}))["result"]
    assert result["success_probability_scope"] == "in_sample"
    oos = result["out_of_sample"]
    assert oos["scenarios"] == 2000 and oos["success_probability"] == pytest.approx(report.oos.pr)
    assert oos["ci95_low"] <= oos["success_probability"] <= oos["ci95_high"] and "meets_threshold" in oos
    assert result["expected_path_return_arithmetic"] == pytest.approx(report.expected_path_return)
    assert result["expected_path_growth_compound"] == pytest.approx(report.expected_path_growth)
    assert result["expected_path_growth_compound"] < result["expected_path_return_arithmetic"]


def test_json_without_validation_has_no_out_of_sample_block(market, artefacts):
    plan, report, _ = artefacts
    from dataclasses import replace
    assert json.loads(build_json(replace(report, oos=None), plan, market, {}))["result"]["out_of_sample"] is None


def test_the_excel_summary_shows_probabilities_and_rates_as_percentages(market, artefacts):
    plan, report, _ = artefacts
    wb = load_workbook(io.BytesIO(build_excel(report, plan, market, {}, None)))
    formats = {row[0].value: row[1].number_format for row in wb["Sintesi"].iter_rows(min_row=2) if row[0].value}
    for key in ("export.pr_in", "export.irr_tail", "export.required", "export.expected_growth", "export.pr_target"):
        assert formats[t(key)] == "0.00%", key
    assert formats[t("export.pr_oos", n="2,000")] == "0.00%" and formats[t("export.horizon")] == "General"


def test_the_goal_column_says_when_the_amounts_are_adjusted_for_inflation(market):
    plan = Plan(w0=60_000, contribution=4_000, horizon=15, goals=((15, 200_000.0),), pr_target=0.9, alpha=0.05,
                inflation=0.02)
    flows = plan.flows()
    problem = GoalProblem(flows, market.draw_scenarios(15, 300, np.random.default_rng(1)), 0.9, 0.05)
    report = build_report(market, flows, problem, run_ga(problem, GAConfig(pop_size=30, patience=10, seed=2)),
                          oos_scenarios=0)
    goals = load_workbook(io.BytesIO(build_excel(report, plan, market, {}, None)))["Obiettivi"]
    assert [c.value for c in goals[1]][1] == t("export.col_goal_adj")
    assert goals.cell(row=2, column=2).value == pytest.approx(200_000 * 1.02**15)


def test_json_does_not_depend_on_the_interface_language(market, artefacts):
    plan, report, _ = artefacts
    docs = []
    for lang in ("it", "en", "zh"):
        with i18n.use(lang):
            payload = json.loads(build_json(report, plan, market, {"seed": 2}))
        payload.pop("generated_at")
        docs.append(payload)
    assert docs[0] == docs[1] == docs[2]
