"""Esportazione dei risultati (Excel nella lingua attiva e JSON riproducibile con chiavi in inglese)."""
from __future__ import annotations

import io
import json
import re
from dataclasses import asdict
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl.styles import Font

from goalga.analytics import OptimizationReport, RobustnessReport
from goalga.i18n import has, t
from goalga.market import Market
from goalga.plan import Plan

from . import charts

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet(name_key: str) -> str:
    return _BAD_SHEET_CHARS.sub("", t(name_key))[:31]


def _setting_label(key: str) -> str:
    label_key = f"export.setting.{key}"
    return t(label_key) if has(label_key) else key


_PERCENT_ROWS = ("export.growth", "export.inflation", "export.pr_target", "export.alpha", "export.pr_in",
                 "export.irr_tail", "export.required", "export.expected_return", "export.expected_growth",
                 "export.ci_low", "export.ci_high", "export.irr_oos")


def _summary_rows(report: OptimizationReport, plan: Plan, market: Market, settings: dict) -> pd.DataFrame:
    rows = [
        (t("export.h.plan"), ""),
        (t("export.w0"), plan.w0),
        (t("export.contribution"), plan.contribution),
        (t("export.horizon"), plan.horizon),
        (t("export.last_contribution"), plan.effective_last_contribution()),
        (t("export.growth"), plan.contribution_growth),
        (t("export.inflation"), plan.inflation),
        (t("export.pr_target"), plan.pr_target),
        (t("export.alpha"), plan.alpha),
        ("", ""),
        (t("export.h.results"), ""),
        (t("export.pr_in"), report.pr),
        (t("export.irr_tail"), report.irr_es),
        (t("export.tail_wealth"), report.tail_wealth),
        (t("export.required"), report.required_return if report.required_return is not None else t("an.na")),
        (t("export.expected_return"), report.expected_path_return),
        (t("export.expected_growth"), report.expected_path_growth),
        (t("export.switches"), report.switches),
    ]
    if report.oos:
        rows += [
            (t("export.pr_oos", n=f"{report.oos.n_scenarios:,}"), report.oos.pr),
            (t("export.ci_low"), report.oos.pr_low),
            (t("export.ci_high"), report.oos.pr_high),
            (t("export.irr_oos"), report.oos.irr_es),
        ]
    start, end = market.meta.get("start"), market.meta.get("end")
    rows += [("", ""), (t("export.h.market"), ""),
             (t("export.assets"), ", ".join(market.names)),
             (t("export.window"), f"{start:%m/%Y} - {end:%m/%Y} ({market.n_months})" if start is not None else ""),
             (t("export.bootstrap"), market.method),
             (t("export.shrinkage"), market.shrinkage),
             (t("export.custom_portfolios"), t("export.yes") if market.custom_weights else t("export.no"))]
    rows += [(_setting_label(k), t(f"risk.{v}.name") if k == "risk_profile" and has(f"risk.{v}.name") else v)
             for k, v in settings.items()]
    return pd.DataFrame(rows, columns=[t("export.col_item"), t("export.col_value")])


def build_excel(report: OptimizationReport, plan: Plan, market: Market, settings: dict,
                robustness: RobustnessReport | None = None) -> bytes:
    ret, vol = market.stats
    glide = report.path_stats.reset_index()
    glide["Portafoglio"] = "P" + glide["Portafoglio"].astype(str)
    glide.columns = [t("export.col_year"), t("export.col_portfolio"), t("export.col_exp_ret"), t("export.col_vol")]
    portfolios = market.weights_frame()
    portfolios.loc[t("export.col_exp_ret")] = ret
    portfolios.loc[t("export.col_vol")] = vol
    portfolios = portfolios.reset_index().rename(columns={"index": t("export.col_asset")})
    assumptions = pd.DataFrame({
        t("export.col_asset"): market.names, t("export.col_hist_ret"): market.mu_hist,
        t("export.col_used_ret"): market.mu, t("export.col_vol"): np.sqrt(np.diag(market.cov)),
    })
    composition = report.composition.reset_index()
    composition.columns = [t("export.col_year"), *composition.columns[1:]]
    fan = report.fan.reset_index()
    fan.columns = [t("export.col_year"), *[f"P{q}" for q in (5, 25, 50, 75, 95)], t("export.col_expected")]
    goal_label = t("export.col_goal_adj") if plan.inflation > 0 else t("export.col_goal")
    goals = report.goals.rename(columns={
        "Anno": t("export.col_year"), "Obiettivo": goal_label,
        "Probabilità cumulata di successo": t("export.col_cum_prob"), "Copertura mediana": t("export.col_cov_median"),
        "Copertura nel 5° percentile": t("export.col_cov_p5")})
    bench = report.benchmarks_validated.copy()
    strategy_labels = [charts.strategy_label(k) for k in bench["Chiave"]]
    bench = bench.drop(columns=["Chiave", "Strategia"]).rename(columns={
        "Probabilità di successo": t("export.col_prob"), "IRR atteso (coda)": t("export.col_irr_tail"),
        "Patrimonio finale medio (coda)": t("export.col_tail_wealth"), "Soglia rispettata": t("export.col_ok"),
        "Score": t("export.col_score")})
    bench.insert(0, t("export.col_strategy"), strategy_labels)
    sheets = {
        _sheet("export.sheet.summary"): (_summary_rows(report, plan, market, settings), set()),
        _sheet("export.sheet.glide"): (glide, {t("export.col_exp_ret"), t("export.col_vol")}),
        _sheet("export.sheet.composition"): (composition, set(composition.columns[1:])),
        _sheet("export.sheet.portfolios"): (portfolios, set()),
        _sheet("export.sheet.assumptions"): (assumptions, {t("export.col_hist_ret"), t("export.col_used_ret"),
                                                           t("export.col_vol")}),
        _sheet("export.sheet.wealth"): (fan, set()),
        _sheet("export.sheet.goals"): (goals, {t("export.col_cum_prob")}),
        _sheet("export.sheet.comparison"): (bench, {t("export.col_prob"), t("export.col_irr_tail")}),
    }
    if robustness is not None:
        rob = pd.DataFrame(robustness.paths,
                           columns=[t("chart.year_n", n=k) for k in range(1, robustness.paths.shape[1] + 1)])
        rob[t("export.col_probability")] = robustness.pr
        rob["IRR"] = robustness.irr_es
        rob = rob.rename_axis(t("export.col_run")).reset_index()
        sheets[_sheet("export.sheet.robustness")] = (rob, {t("export.col_probability"), "IRR"})
    summary_name = _sheet("export.sheet.summary")

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, (frame, pct_columns) in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for idx, col in enumerate(frame.columns, start=1):
                width = max(len(str(col)), *(len(str(v)) for v in frame[col].head(60))) + 2
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(width, 10), 60)
                if name != summary_name and col in pct_columns and pd.api.types.is_numeric_dtype(frame[col]):
                    for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                        row[0].number_format = "0.00%"
            if name == summary_name:                          # righe di percentuali e probabilità: mostrate come tali
                percent_labels = {t(k) for k in _PERCENT_ROWS} | {t("export.pr_oos", n=f"{report.oos.n_scenarios:,}")
                                                                  if report.oos else ""}
                for row in ws.iter_rows(min_row=2, max_col=2):
                    if row[0].value in percent_labels and isinstance(row[1].value, (int, float)):
                        row[1].number_format = "0.00%"
    return buffer.getvalue()


def build_json(report: OptimizationReport, plan: Plan, market: Market, settings: dict) -> str:
    """Dati completi in JSON (chiavi in inglese, indipendenti dalla lingua dell'interfaccia)."""
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "plan": {**asdict(plan), "goals": [list(g) for g in plan.goals]},
        "market": {
            "assets": market.names, "expected_returns_used": market.mu.tolist(),
            "expected_returns_historical": market.mu_hist.tolist(), "bootstrap_method": market.method,
            "shrinkage": market.shrinkage, "portfolio_weights": market.weights.T.tolist(),
            "custom_portfolios": market.custom_weights,
        },
        "settings": settings,
        "result": {
            "glide_path": report.path.tolist(), "success_probability": report.pr, "tail_irr": report.irr_es,
            "tail_wealth": report.tail_wealth, "required_return": report.required_return,
            "expected_path_return_arithmetic": report.expected_path_return,
            "expected_path_growth_compound": report.expected_path_growth,
            "generations": report.ga.generations, "stop_reason": report.ga.stopped_by,
            "success_probability_scope": "in_sample",
            "out_of_sample": None if report.oos is None else {
                "scenarios": report.oos.n_scenarios, "success_probability": report.oos.pr,
                "ci95_low": report.oos.pr_low, "ci95_high": report.oos.pr_high, "tail_irr": report.oos.irr_es,
                "meets_threshold": bool(report.oos.feasible)},
        },
    }
    return json.dumps(payload, ensure_ascii=False, indent=2, default=lambda o: o.item() if hasattr(o, "item") else str(o))
